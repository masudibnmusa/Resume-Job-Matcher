import logging
from pathlib import Path

from app.llm.llm_client import LLMClient
from app.llm.llm_client import LLMError
from app.models import Education, Experience, Resume
from app.suggestions.prompt_templates import SYSTEM_EXTRACT, resume_extraction_prompt
from app.parsing.section_extractor import rule_based_resume

log = logging.getLogger(__name__)


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume format: {suffix} (use PDF, DOCX, or TXT)")


def _resume_from_dict(d: dict, raw_text: str) -> Resume:
    experience = []
    for e in d.get("experience") or []:
        if not isinstance(e, dict):
            continue
        experience.append(
            Experience(
                title=_s(e.get("title")),
                company=_s(e.get("company")),
                start=_s(e.get("start")),
                end=_s(e.get("end")),
                bullets=[_s(b) for b in (e.get("bullets") or []) if _s(b)],
            )
        )
    education = []
    for ed in d.get("education") or []:
        if not isinstance(ed, dict):
            continue
        education.append(
            Education(
                degree=_s(ed.get("degree")),
                field=_s(ed.get("field")),
                institution=_s(ed.get("institution")),
                year=_s(ed.get("year")),
            )
        )
    return Resume(
        name=_s(d.get("name")),
        summary=_s(d.get("summary")),
        skills=[_s(s) for s in (d.get("skills") or []) if _s(s)],
        experience=experience,
        education=education,
        raw_text=raw_text,
    )


def parse_resume(path: str | Path, llm: LLMClient | None = None) -> Resume:
    text = extract_text(path)
    if not text.strip():
        raise ValueError("No text found in the resume. Scanned/image PDFs need OCR, which is not supported.")
    if llm is not None:
        try:
            data = llm.complete_json(resume_extraction_prompt(text), system=SYSTEM_EXTRACT)
            if isinstance(data, dict):
                return _resume_from_dict(data, text)
        except (LLMError, ValueError) as e:
            log.warning("LLM resume extraction failed (%s); using rule-based parser.", e)
    return rule_based_resume(text)